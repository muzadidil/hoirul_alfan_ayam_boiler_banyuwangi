#!/usr/bin/env python3
"""Impor mutasi BCA dari Excel -> bersihkan -> periksa -> (unggah ke Firestore).

Contoh:
  python alat/impor.py --cek
  FIRESTORE_EMULATOR_HOST=localhost:8080 python alat/impor.py --unggah
  python alat/impor.py --unggah --proyek zasha-mutasi
  python alat/impor.py --unggah --bulan 2023-03 --timpa

Tanpa --timpa, transaksi yang sudah ada di Firestore DILEWATI supaya perubahan
kategori/entitas dari aplikasi tidak tertimpa. Ringkasan bulan selalu dihitung
dari data di Firestore setelah penulisan.
"""
import argparse
import csv
import datetime as dt
import glob
import os
import secrets
import sys
from collections import Counter

import bersih as B
import koneksi

FOLDER_REPO = koneksi.FOLDER_REPO

# Angka pemeriksaan hasil audit data 2023 (jumlah baris, bukan nominal).
ANGKA_CEK = {
    2023: {
        'jumlahPerBulan': {'2023-01': 679, '2023-02': 548, '2023-03': 561, '2023-04': 368, '2023-05': 461,
                           '2023-06': 595, '2023-07': 620, '2023-08': 663, '2023-09': 661, '2023-10': 623,
                           '2023-11': 672, '2023-12': 588},
        'kategoriAsal': {'PRODUKSI': 2484, 'PRIBADI': 2341, None: 1764, 'EXSPEDISI': 283, 'V_SALES': 73,
                         'BANK': 61, 'NURUL_AINI': 30, 'OPERASIONAL': 3},
        # Celah saldo yang sudah diketahui: mutasi MEI & AGU tidak memuat baris bunga/pajak bunga akhir bulan.
        'celahSaldo': {('2023-05', '2023-06'), ('2023-08', '2023-09')},
        # J1 tab kategori melewatkan transaksi pertama (rumus mulai D2).
        'pengecualianJ1': {('2023-01', 'PRODUKSI'), ('2023-02', 'PRODUKSI'), ('2023-02', 'PRIBADI')},
    },
}
CELAH_MAKS_SEN = 1000000  # celah bunga-pajak wajar < Rp 10.000
TOLERANSI_SEN = 100       # Rp 1 untuk nilai cache Google Sheets


# ---------------------------------------------------------------- baca Excel

def baca_folder(folder):
    """Cari tab master <MMM>_<TAHUN> di semua .xlsx, tab kategori (J1), dan BANK_GLOBAL."""
    import openpyxl

    master, bank_global = [], {}
    berkas_list = sorted(glob.glob(os.path.join(folder, '*.xlsx')))
    berkas_list = [f for f in berkas_list if not os.path.basename(f).startswith('~$')]
    for jalur in berkas_list:
        nama = os.path.basename(jalur)
        wb = openpyxl.load_workbook(jalur, data_only=True)
        for tab in wb.sheetnames:
            m = B.POLA_TAB.match(tab)
            if m:
                ws = wb[tab]
                baris = [(n, list(r)) for n, r in
                         enumerate(ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=9, values_only=True), 1)]
                j1 = {}
                for lain in wb.sheetnames:
                    kode_kat = B.normal_kategori(lain)
                    if lain != tab and kode_kat in B.KATEGORI:
                        j1[kode_kat] = wb[lain]['J1'].value
                master.append({'tahun': int(m.group(2)), 'bulan': B.KODE_BULAN.index(m.group(1)) + 1,
                               'tab': tab, 'berkas': nama, 'baris': baris, 'j1': j1})
            elif tab.upper() == 'TOTAL':
                bank_global.update(_baca_bank_global(wb[tab], nama))
        wb.close()
    master.sort(key=lambda x: (x['tahun'], x['bulan']))
    return master, bank_global, [os.path.basename(f) for f in berkas_list]


def _baca_bank_global(ws, nama):
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return {}
    kepala = [B.rapikan_spasi(v).upper() for v in rows[0]]
    if 'KREDIT' not in kepala:
        return {}
    k = kepala.index('KREDIT')
    hasil = {}
    for r in rows[1:]:
        nama_bln = B.rapikan_spasi(r[0]).upper() if r and r[0] else ''
        if nama_bln in [n.upper() for n in B.NAMA_BULAN]:
            hasil[[n.upper() for n in B.NAMA_BULAN].index(nama_bln) + 1] = {'kredit': r[k], 'berkas': nama}
    return hasil


# ---------------------------------------------------------------- bersihkan

def proses(master, pakai_kategori_asal=False, penyesuaian=False):
    daftar_bulan = []
    for m in master:
        bln = B.bersihkan_bulan(m['baris'], m['tahun'], m['bulan'], m['tab'], m['berkas'])
        bln['j1'] = m['j1']
        daftar_bulan.append(bln)
    semua = [t for b in daftar_bulan for t in b['transaksi']]
    peta, rincian_peta = B.buat_peta_saran(semua)
    for t in semua:
        B.terapkan_kategori(t, pakai_kategori_asal, peta)
    for bln in daftar_bulan:
        ragu = sum(1 for t in bln['transaksi'] if 'kategori_diragukan' in t['perluCek'])
        if ragu:
            bln['catatan'].append('Kategori kolom I bulan ini tidak dipercaya (tersalin per posisi baris): '
                                  '%d transaksi diimpor sebagai KOSONG, nilai lama di kategoriAsal' % ragu)
    for a, b in zip(daftar_bulan, daftar_bulan[1:]):
        if a['saldoAkhirSen'] is None or b['saldoAwalSen'] is None:
            continue
        celah = b['saldoAwalSen'] - a['saldoAkhirSen']
        a['celahKeBerikutSen'] = celah
        if celah and _berurutan(a, b):
            a['catatan'].append('Saldo akhir berbeda %s dengan saldo awal %s (baris bunga/pajak bunga akhir bulan '
                                'kemungkinan tidak tersalin)' % (B.rupiah(celah), b['kode']))
            if penyesuaian:
                t = B.transaksi_penyesuaian(a, celah)
                a['transaksi'].append(t)
                a['saldoAkhirSen'] = t['saldoSen']
                a['catatan'].append('Baris PENYESUAIAN sintetis %s ditambahkan (%s)' % (t['id'], B.rupiah(celah)))
    return daftar_bulan, rincian_peta


def _berurutan(a, b):
    return (a['tahun'] * 12 + a['bulan']) + 1 == b['tahun'] * 12 + b['bulan']


# ---------------------------------------------------------------- pemeriksaan

class Pemeriksaan:
    def __init__(self):
        self.hasil = []

    def tambah(self, nama, lulus, rincian=()):
        self.hasil.append((nama, bool(lulus), list(rincian)))

    @property
    def lulus(self):
        return all(l for _, l, _ in self.hasil)


def periksa(daftar_bulan, bank_global, penyesuaian=False):
    P = Pemeriksaan()
    semua = [t for b in daftar_bulan for t in b['transaksi']]
    nyata = [t for t in semua if 'sintetis' not in t['perluCek']]
    tahun_ada = sorted({b['tahun'] for b in daftar_bulan})

    galat = [('%s: %s' % (b['tab'], g)) for b in daftar_bulan for g in b['galat']]
    P.tambah('Tidak ada galat saat membaca baris', not galat and daftar_bulan, galat or ['%d tab master' % len(daftar_bulan)])

    for tahun in tahun_ada:
        cek = ANGKA_CEK.get(tahun)
        if not cek:
            continue
        per_bulan = Counter(t['bulan'] for t in nyata if t['tahun'] == tahun)
        rinci = ['%s: %d (harus %d)%s' % (k, per_bulan.get(k, 0), v, '' if per_bulan.get(k, 0) == v else '  <-- BEDA')
                 for k, v in sorted(cek['jumlahPerBulan'].items())]
        total = sum(per_bulan.values())
        harus = sum(cek['jumlahPerBulan'].values())
        P.tambah('Jumlah transaksi %d = %s' % (tahun, '{:,}'.format(harus).replace(',', '.')),
                 dict(per_bulan) == cek['jumlahPerBulan'],
                 ['Total %s' % '{:,}'.format(total).replace(',', '.')] + rinci)

    ids = [t['id'] for t in semua]
    ganda = [i for i, n in Counter(ids).items() if n > 1]
    P.tambah('ID transaksi unik', not ganda, ganda[:20] or ['%d ID unik' % len(ids)])

    masalah = ['%s tanpa tanggal' % t['id'] for t in semua if not t.get('tanggal')]
    masalah += ['%s bertanggal %s di luar bulannya' % (t['id'], t['tanggal']) for t in semua
                if t.get('tanggal') and t['tanggal'][:7] != t['bulan']]
    masalah += ['%s (%s) > %s (%s)' % (x['id'], x['tanggal'], y['id'], y['tanggal'])
                for x, y in zip(semua, semua[1:]) if x.get('tanggal') and y.get('tanggal') and x['tanggal'] > y['tanggal']]
    rentang = '%s s.d. %s' % (semua[0]['tanggal'], semua[-1]['tanggal']) if semua else 'tidak ada transaksi'
    P.tambah('Semua transaksi bertanggal, di bulannya, dan urut naik sepanjang tahun', semua and not masalah,
             masalah[:20] or [rentang])

    salah_sisi = [t['id'] for t in semua if (t['debetSen'] > 0) == (t['kreditSen'] > 0)
                  or t['debetSen'] < 0 or t['kreditSen'] < 0
                  or t['jenis'] != ('DB' if t['debetSen'] > 0 else 'CR')]
    P.tambah('Tepat satu dari debet/kredit terisi, jenis sesuai nominal', not salah_sisi,
             salah_sisi[:20] or ['%d DB, %d CR' % (sum(t['jenis'] == 'DB' for t in semua),
                                                   sum(t['jenis'] == 'CR' for t in semua))])

    # Kredit sebelum koreksi BUNGA vs BANK_GLOBAL kolom KREDIT
    rinci, ok = [], True
    for b in daftar_bulan:
        kredit_awal = sum(t.get('_kreditSebelumBungaSen', t['kreditSen']) for t in b['transaksi']
                          if 'sintetis' not in t['perluCek'])
        bg = bank_global.get(b['bulan']) if b['tahun'] == 2023 else None
        bg_sen = B.ke_sen(bg['kredit']) if bg else None
        e2_sen = b.get('kreditBarisSaldoAwalSen')
        if e2_sen is not None and abs(kredit_awal - e2_sen) > TOLERANSI_SEN:
            ok = False
            rinci.append('%s: kredit %s | E%d (SUM kredit di baris SALDO AWAL) %s  <-- BEDA' % (
                b['kode'], B.rupiah(kredit_awal), b['barisSaldoAwal'], B.rupiah(e2_sen)))
        if bg_sen is None:
            rinci.append('%s: kredit %s | BANK_GLOBAL kosong, dilewati' % (b['kode'], B.rupiah(kredit_awal)))
            continue
        beda = kredit_awal - bg_sen
        cocok = abs(beda) <= TOLERANSI_SEN
        ok &= cocok
        rinci.append('%s: kredit %s | BANK_GLOBAL %s | selisih %s%s' % (
            b['kode'], B.rupiah(kredit_awal), B.rupiah(bg_sen), B.rupiah(beda), '' if cocok else '  <-- BEDA'))
    P.tambah('Kredit per bulan (sebelum koreksi BUNGA) = BANK_GLOBAL KREDIT (selisih <= Rp 1)',
             ok and bank_global, rinci if bank_global else ['BANK_GLOBAL.xlsx tidak ditemukan'])

    # J1 tab kategori (hanya debet; QUERY mengabaikan debet berupa teks)
    rinci, ok = [], True
    for b in daftar_bulan:
        cek = ANGKA_CEK.get(b['tahun'], {})
        for kat, j1 in sorted(b['j1'].items()):
            cocok_baris = [t for t in b['transaksi'] if t.get('_kategoriMentah', t['kategoriAsal'] or '') ==
                           ('' if kat == 'KOSONG' else kat) and t.get('_dNumerikSen')]
            hitung = sum(t['_dNumerikSen'] for t in cocok_baris)
            j1_sen = B.ke_sen(j1) if isinstance(j1, (int, float)) else 0
            beda = j1_sen - hitung
            if abs(beda) <= TOLERANSI_SEN:
                continue
            pertama = cocok_baris[0]['_dNumerikSen'] if cocok_baris else None
            dikenal = (b['kode'], kat) in cek.get('pengecualianJ1', set()) and pertama is not None \
                and abs(beda + pertama) <= TOLERANSI_SEN
            ok &= dikenal
            rinci.append('%s %s: J1 %s | hitung %s | selisih %s%s' % (
                b['kode'], kat, B.rupiah(j1_sen), B.rupiah(hitung), B.rupiah(beda),
                ' (pengecualian diketahui: J1 melewatkan transaksi pertama %s)' % cocok_baris[0]['id']
                if dikenal else '  <-- BEDA'))
    P.tambah('Debet per kategori (kategoriAsal) = J1 tab kategori, kecuali pengecualian yang diketahui', ok,
             rinci or ['semua cocok'])

    # Saldo: kolom H dan sambungan antarbulan
    beda_h = [t['id'] for t in semua if 'saldo_beda' in t['perluCek']]
    beda_h += ['%s baris SALDO AWAL: H %s != G %s' % (b['kode'], B.rupiah(b['saldoBankAwalSen']), B.rupiah(b['saldoAwalSen']))
               for b in daftar_bulan if b.get('saldoBankAwalSen') is not None and b['saldoAwalSen'] is not None
               and abs(b['saldoBankAwalSen'] - b['saldoAwalSen']) > 1]
    sel_h = sum(b['kontrol']['selH'] for b in daftar_bulan)
    nyasar = ['%s r%d' % (b['kode'], n) for b in daftar_bulan for n in b.get('hNyasar', [])]
    P.tambah('Saldo hitung = kolom H (saldo cetak bank)', not beda_h,
             beda_h[:20] or ['%d sel H transaksi cocok (+%d sel H di baris SALDO AWAL); '
                             '%d sel H berisi nominal (bukan saldo) diabaikan: %s'
                             % (sel_h - len(nyasar), sum(b.get('saldoBankAwalSen') is not None for b in daftar_bulan),
                                len(nyasar), ', '.join(nyasar) or '-')])
    rinci, ok = [], True
    for a, b in zip(daftar_bulan, daftar_bulan[1:]):
        if not _berurutan(a, b) or a['saldoAkhirSen'] is None or b['saldoAwalSen'] is None:
            continue
        celah = b['saldoAwalSen'] - a['saldoAkhirSen']
        dikenal = (a['kode'], b['kode']) in ANGKA_CEK.get(a['tahun'], {}).get('celahSaldo', set())
        if celah == 0:
            rinci.append('%s -> %s: tersambung' % (a['kode'], b['kode']))
            continue
        lulus = dikenal and 0 < abs(celah) < CELAH_MAKS_SEN
        ok &= lulus
        rinci.append('%s -> %s: celah %s%s' % (a['kode'], b['kode'], B.rupiah(celah),
                                                ' (pengecualian diketahui)' if lulus else '  <-- BEDA'))
    if penyesuaian:
        ok &= not any('celah' in r for r in rinci)
    P.tambah('Saldo akhir bulan = saldo awal bulan berikutnya, kecuali celah yang diketahui', ok, rinci)

    # Baris penutup bank (MUTASI DB / MUTASI CR)
    rinci, ok, ada = [], True, False
    for b in daftar_bulan:
        nyata_b = [t for t in b['transaksi'] if 'sintetis' not in t['perluCek']]
        for p in b['penutup']:
            if p['jenis'] not in ('MUTASI DB', 'MUTASI CR'):
                continue
            ada = True
            jenis = p['jenis'][-2:]
            try:
                nominal, jumlah = B.ke_sen(p['d']), int(B.parse_angka(p['e']) or 0)
            except ValueError as err:
                ok = False
                rinci.append('%s %s: tidak terbaca (%s)' % (b['kode'], p['jenis'], err))
                continue
            hitung_n = sum(1 for t in nyata_b if t['jenis'] == jenis)
            hitung_sen = sum(t['debetSen' if jenis == 'DB' else 'kreditSen'] for t in nyata_b)
            cocok = nominal == hitung_sen and jumlah == hitung_n
            ok &= cocok
            rinci.append('%s %s: bank %s / %d trx | hitung %s / %d trx%s' % (
                b['kode'], p['jenis'], B.rupiah(nominal), jumlah, B.rupiah(hitung_sen), hitung_n,
                '' if cocok else '  <-- BEDA'))
    harus_ada = any(b['kode'] == '2023-09' for b in daftar_bulan)
    P.tambah('Total debet & kredit = baris penutup bank (MUTASI DB / MUTASI CR)', ok and (ada or not harus_ada),
             rinci or ['tidak ada baris penutup'])

    for tahun in tahun_ada:
        cek = ANGKA_CEK.get(tahun)
        if not cek:
            continue
        c = Counter(t['kategoriAsal'] for t in nyata if t['tahun'] == tahun)
        rinci = ['%s: %d (harus %d)%s' % (k or '(kosong)', c.get(k, 0), v, '' if c.get(k, 0) == v else '  <-- BEDA')
                 for k, v in cek['kategoriAsal'].items()]
        lain = {k: v for k, v in c.items() if k not in cek['kategoriAsal']}
        if lain:
            rinci.append('kategori lain: %s  <-- BEDA' % lain)
        P.tambah('Jumlah per kategoriAsal %d' % tahun, dict(c) == cek['kategoriAsal'], rinci)

    tak_dikenal = sorted({t['kategoriAsal'] for t in semua if t['kategoriAsal'] and t['kategoriAsal'] not in B.KATEGORI})
    P.tambah('Semua kategoriAsal dikenal', not tak_dikenal, tak_dikenal or ['ok'])
    return P


# ---------------------------------------------------------------- keluaran

def tulis_keluaran(daftar_bulan, rincian_peta, P, folder, opsi):
    os.makedirs(folder, exist_ok=True)
    tulisan = []
    for tahun in sorted({b['tahun'] for b in daftar_bulan}):
        bulan_th = [b for b in daftar_bulan if b['tahun'] == tahun]
        jalur = os.path.join(folder, 'transaksi_%d.csv' % tahun)
        kolom = ['id', 'bulan', 'tanggal', 'urutan', 'jenis', 'debetSen', 'kreditSen', 'saldoSen', 'saldoBankSen',
                 'kategori', 'kategoriAsal', 'saran', 'entitas', 'perluCek', 'ketPendek', 'ketLengkap', 'ketAsli',
                 'file', 'tab', 'baris']
        with open(jalur, 'w', newline='', encoding='utf-8-sig') as f:
            w = csv.writer(f)
            w.writerow(kolom)
            for b in bulan_th:
                for t in b['transaksi']:
                    w.writerow([t['id'], t['bulan'], t['tanggal'], t['urutan'], t['jenis'], t['debetSen'],
                                t['kreditSen'], t['saldoSen'], _kosong(t['saldoBankSen']), t['kategori'],
                                _kosong(t['kategoriAsal']), _kosong(t['saran']), t['entitas'], '|'.join(t['perluCek']),
                                t['ketPendek'], t['ketLengkap'], _kosong(t['ketAsli']), t['sumber']['file'],
                                t['sumber']['tab'], _kosong(t['sumber']['baris'])])
        tulisan.append(jalur)

        jalur = os.path.join(folder, 'ringkasan_%d.csv' % tahun)
        with open(jalur, 'w', newline='', encoding='utf-8-sig') as f:
            w = csv.writer(f)
            w.writerow(['bulan', 'entitas', 'kategori', 'debetSen', 'kreditSen', 'jumlah', 'debetRp', 'kreditRp',
                        'saldoAwalSen', 'saldoAkhirSen'])
            for b in bulan_th:
                r = B.hitung_ringkasan(b['transaksi'])
                per_ek = Counter((t['entitas'], t['kategori']) for t in b['transaksi'])
                awal, akhir = b['saldoAwalSen'], b['saldoAwalSen'] + r['totalKreditSen'] - r['totalDebetSen']
                baris = [('SEMUA', 'SEMUA', r['totalDebetSen'], r['totalKreditSen'], r['jumlahTransaksi'], awal, akhir)]
                baris += [('SEMUA', k, r['debetSenPerKategori'][k], r['kreditSenPerKategori'][k],
                           r['jumlahPerKategori'][k], '', '') for k in B.KATEGORI]
                baris += [(e, 'SEMUA', r['debetSenPerEntitas'][e], r['kreditSenPerEntitas'][e],
                           r['jumlahPerEntitas'][e], '', '') for e in B.ENTITAS]
                baris += [(e, k, r['debetSenPerEntitasKategori'][e][k], r['kreditSenPerEntitasKategori'][e][k],
                           per_ek[(e, k)], '', '') for e in B.ENTITAS for k in B.KATEGORI]
                for e, k, d, c, n, sa, sk in baris:
                    w.writerow([b['kode'], e, k, d, c, n, '%.2f' % (d / 100), '%.2f' % (c / 100), sa, sk])
        tulisan.append(jalur)

    jalur = os.path.join(folder, 'peta_saran.csv')
    with open(jalur, 'w', newline='', encoding='utf-8-sig') as f:
        w = csv.writer(f)
        w.writerow(['jenis', 'pihak', 'kategori', 'jumlahKategori', 'jumlahBerkategori', 'jumlahSemua', 'persen',
                    'dipakai', 'sebaran'])
        for r in rincian_peta:
            w.writerow([r['jenis'], r['pihak'], r['kategori'], r['jumlahKategori'], r['jumlahBerkategori'],
                        r['jumlahSemua'], r['persen'], 'ya' if r['dipakai'] else 'tidak',
                        ' '.join('%s=%d' % kv for kv in sorted(r['sebaran'].items()))])
    tulisan.append(jalur)

    jalur = os.path.join(folder, 'laporan_cek.txt')
    with open(jalur, 'w', encoding='utf-8') as f:
        f.write(laporan_teks(daftar_bulan, rincian_peta, P, opsi))
    tulisan.append(jalur)
    return tulisan


def _kosong(v):
    return '' if v is None else v


def laporan_teks(daftar_bulan, rincian_peta, P, opsi):
    semua = [t for b in daftar_bulan for t in b['transaksi']]
    out = ['LAPORAN PEMERIKSAAN IMPOR ZASHA MUTASI',
           'Dibuat: %s  |  versi alat %s' % (dt.datetime.now().strftime('%Y-%m-%d %H:%M'), B.VERSI),
           'Opsi: %s' % ', '.join('%s=%s' % kv for kv in sorted(opsi.items())), '',
           'HASIL: %s' % ('LULUS' if P.lulus else 'GAGAL'), '']
    for nama, lulus, rinci in P.hasil:
        out.append('[%s] %s' % ('LULUS' if lulus else 'GAGAL', nama))
        out += ['        ' + r for r in rinci]
    out += ['', 'RINGKASAN PER BULAN',
            '%-8s %-14s %5s %22s %22s %22s %22s %18s' % ('bulan', 'tab', 'trx', 'saldo awal', 'debet', 'kredit',
                                                         'saldo akhir', 'celah ke berikut')]
    for b in daftar_bulan:
        d = sum(t['debetSen'] for t in b['transaksi'])
        c = sum(t['kreditSen'] for t in b['transaksi'])
        out.append('%-8s %-14s %5d %22s %22s %22s %22s %18s' % (
            b['kode'], b['tab'], len(b['transaksi']), B.rupiah(b['saldoAwalSen']), B.rupiah(d), B.rupiah(c),
            B.rupiah(b['saldoAkhirSen']), B.rupiah(b.get('celahKeBerikutSen')) if 'celahKeBerikutSen' in b else '-'))
    out += ['', 'PERBAIKAN OTOMATIS PER BULAN']
    kunci = [('barisKosong', 'baris kosong'), ('ketGmt', 'teks GMT'), ('ketGanda', 'ket. ganda'),
             ('labelSaldo', 'label C'), ('bDariC', 'B dari C'), ('angkaTeks', 'angka teks'),
             ('bungaDipindah', 'BUNGA'), ('hNyasar', 'H nyasar')]
    out.append('%-8s ' % 'bulan' + ' '.join('%12s' % k[1] for k in kunci) + ' %12s %12s' % ('tgl ditukar', 'kat ragu'))
    for b in daftar_bulan:
        f = Counter(x for t in b['transaksi'] for x in t['perluCek'])
        out.append('%-8s ' % b['kode'] + ' '.join('%12d' % b['kontrol'][k[0]] for k in kunci)
                   + ' %12d %12d' % (f['tanggal_ditukar'], f['kategori_diragukan']))
    out += ['', 'TANDA perluCek (jumlah transaksi)']
    f = Counter(x for t in semua for x in t['perluCek'])
    out += ['  %-24s %d' % (k, f[k]) for k in B.PERLU_CEK]
    out += ['', 'KATEGORI AWAL & ENTITAS AWAL']
    out.append('  kategori: %s' % ', '.join('%s %d' % (k, n) for k, n in
                                             sorted(Counter(t['kategori'] for t in semua).items())))
    out.append('  entitas : %s' % ', '.join('%s %d' % (k, n) for k, n in
                                             sorted(Counter(t['entitas'] for t in semua).items())))
    ber_saran = [t for t in semua if t['saran']]
    out += ['', 'SARAN KATEGORI',
            '  peta: %d pihak dipakai dari %d pihak berkategori di bulan terpercaya (%s)' % (
                sum(r['dipakai'] for r in rincian_peta), len(rincian_peta), ', '.join(sorted(B.BULAN_TERPERCAYA))),
            '  transaksi dengan saran: %d dari %d; saran berbeda dari kategori sekarang: %d' % (
                len(ber_saran), len(semua), sum(1 for t in ber_saran if t['saran'] != t['kategori']))]
    for b in daftar_bulan:
        n = sum(1 for t in b['transaksi'] if t['saran'])
        out.append('  %s: %d/%d bersaran' % (b['kode'], n, len(b['transaksi'])))
    out += ['', 'CATATAN PER BULAN (juga disimpan di bulan/{kode}.kualitas.catatan)']
    for b in daftar_bulan:
        out.append('  %s (%s / %s)' % (b['kode'], b['berkas'], b['tab']))
        out += ['    - ' + c for c in b['catatan']] or ['    -']
    out += ['', 'BARIS BERTANDA saldo_beda / sintetis / tanggal_ditukar']
    for t in semua:
        tanda = [x for x in t['perluCek'] if x in ('saldo_beda', 'sintetis', 'tanggal_ditukar')]
        if tanda:
            out.append('  %s  %s  baris %s  %s' % (t['id'], t['tanggal'], t['sumber']['baris'], ','.join(tanda)))
    out.append('')
    return '\n'.join(out)


# ---------------------------------------------------------------- unggah

def unggah(daftar_bulan, args, berkas):
    from google.cloud.firestore import SERVER_TIMESTAMP

    db, tujuan = koneksi.firestore(args.proyek)
    print('Menulis ke %s' % tujuan)
    target = [b for b in daftar_bulan if not args.bulan or b['kode'] in args.bulan]
    total_tulis = total_lewat = 0
    for b in target:
        ref_bulan = db.collection('bulan').document(b['kode'])
        col = ref_bulan.collection('transaksi')
        ada = {d.id: d.to_dict() for d in col.stream()}
        penulis = koneksi.PenulisBatch(db)
        lewat = 0
        for t in b['transaksi']:
            if t['id'] in ada and not args.timpa:
                lewat += 1
                continue
            penulis.set(col.document(t['id']), B.untuk_unggah(t))
        penulis.selesai()
        if penulis.total:
            dokumen = [d.to_dict() for d in col.stream()]
        else:
            dokumen = list(ada.values())
        ringkasan = B.hitung_ringkasan(dokumen)
        dok = B.dokumen_bulan(b['kode'], b['saldoAwalSen'], ringkasan, B.hitung_kualitas(dokumen, b['catatan']),
                              {'file': b['berkas'], 'tab': b['tab']})
        dok['diperbaruiOleh'] = None
        dok['diperbaruiPada'] = SERVER_TIMESTAMP
        ref_bulan.set(dok)
        total_tulis += penulis.total
        total_lewat += lewat
        print('  %s: %4d ditulis, %4d dilewati, %4d transaksi di Firestore' % (
            b['kode'], penulis.total, lewat, len(dokumen)))

    tulis_pengaturan(db)
    id_impor = '%s-%s' % (dt.datetime.now(dt.timezone.utc).strftime('%Y%m%d-%H%M%S'), secrets.token_hex(2))
    db.collection('impor').document(id_impor).set({
        'waktu': SERVER_TIMESTAMP,
        'berkas': berkas,
        'ditulis': total_tulis,
        'dilewati': total_lewat,
        'versi': B.VERSI,
        'opsi': {'bulan': [b['kode'] for b in target], 'timpa': bool(args.timpa),
                 'pakaiKategoriAsal': bool(args.pakai_kategori_asal), 'penyesuaian': bool(args.penyesuaian)},
    })
    print('Selesai: %d ditulis, %d dilewati, %d bulan diperbarui, log impor/%s'
          % (total_tulis, total_lewat, len(target), id_impor))


def tulis_pengaturan(db):
    """Tulis pengaturan/kategori & pengaturan/entitas. Label/ikon/warna yang diubah pemilik dipertahankan."""
    for nama, kode, detail in (('kategori', B.KATEGORI, B.KATEGORI_DETAIL), ('entitas', B.ENTITAS, B.ENTITAS_DETAIL)):
        ref = db.collection('pengaturan').document(nama)
        lama = ref.get()
        baru = {k: dict(v) for k, v in detail.items()}
        if lama.exists:
            for k, v in ((lama.to_dict() or {}).get('detail') or {}).items():
                if k in baru and isinstance(v, dict):
                    baru[k].update({kk: vv for kk, vv in v.items() if kk != 'urutan'})
        ref.set({'kode': list(kode), 'detail': baru})


# ---------------------------------------------------------------- utama

def main(argv=None):
    p = argparse.ArgumentParser(description='Impor mutasi BCA (Excel) ke Firestore.')
    p.add_argument('--folder', default=FOLDER_REPO, help='folder berisi file .xlsx (bawaan: root repo)')
    p.add_argument('--keluaran', default=os.path.join(FOLDER_REPO, 'keluaran'), help='folder hasil --cek')
    p.add_argument('--cek', action='store_true', help='bersihkan & periksa saja, tulis keluaran/*.csv')
    p.add_argument('--unggah', action='store_true', help='tulis ke Firestore (emulator bila FIRESTORE_EMULATOR_HOST)')
    p.add_argument('--proyek', help='ID proyek Firebase (wajib bila bukan emulator)')
    p.add_argument('--bulan', action='append',
                   help='batasi unggah ke bulan tertentu, mis. 2023-03 (boleh berulang); pemeriksaan tetap setahun')
    p.add_argument('--timpa', action='store_true', help='timpa transaksi yang sudah ada (perubahan dari aplikasi hilang)')
    p.add_argument('--pakai-kategori-asal', action='store_true',
                   help='pakai kategori kolom I juga untuk bulan yang diragukan')
    p.add_argument('--penyesuaian', action='store_true',
                   help='tambah baris PENYESUAIAN sintetis untuk menutup celah saldo antarbulan')
    args = p.parse_args(argv)
    if not (args.cek or args.unggah):
        p.error('pilih --cek dan/atau --unggah')
    if args.cek and koneksi.di_dalam_repo(args.keluaran) and \
            os.path.relpath(os.path.realpath(args.keluaran), FOLDER_REPO).split(os.sep)[0] != 'keluaran':
        p.error('--keluaran di dalam repo hanya boleh folder keluaran/ (di-.gitignore)')
    if args.unggah and not koneksi.emulator('firestore') and not args.proyek:
        p.error('--proyek <id> wajib diisi bila tidak memakai emulator (FIRESTORE_EMULATOR_HOST kosong)')

    master, bank_global, berkas = baca_folder(args.folder)
    if not master:
        sys.exit('Galat: tidak ada tab <MMM>_<TAHUN> di %s' % args.folder)
    if args.bulan:
        ada = {'%04d-%02d' % (m['tahun'], m['bulan']) for m in master}
        salah = [k for k in args.bulan if k not in ada]
        if salah:
            p.error('bulan tidak ada di Excel: %s' % ', '.join(salah))
    daftar_bulan, rincian_peta = proses(master, args.pakai_kategori_asal, args.penyesuaian)
    P = periksa(daftar_bulan, bank_global, args.penyesuaian)

    for nama, lulus, rinci in P.hasil:
        print('[%s] %s' % ('LULUS' if lulus else 'GAGAL', nama))
        if not lulus:
            print('\n'.join('        ' + r for r in rinci))
    print('Pemeriksaan: %s (%d transaksi, %d bulan)' % (
        'LULUS' if P.lulus else 'GAGAL', sum(len(b['transaksi']) for b in daftar_bulan), len(daftar_bulan)))

    if args.cek:
        opsi = {'pakaiKategoriAsal': args.pakai_kategori_asal, 'penyesuaian': args.penyesuaian}
        for jalur in tulis_keluaran(daftar_bulan, rincian_peta, P, args.keluaran, opsi):
            print('Ditulis: %s' % jalur)
    if not P.lulus:
        if args.unggah:
            print('Unggah DIBATALKAN karena pemeriksaan gagal.')
        return 1
    if args.unggah:
        try:
            unggah(daftar_bulan, args, berkas)
        except Exception as err:
            print('Galat saat menulis ke Firestore: %s: %s' % (type(err).__name__, err))
            print('Aman untuk dijalankan ulang: transaksi yang sudah tertulis akan dilewati, ringkasan dihitung ulang.')
            return 1
    return 0


if __name__ == '__main__':
    sys.exit(main())
